#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel历史数据记录模块
"""

import os
import time
import threading
from datetime import datetime
from typing import Dict, Any, Optional
import queue

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    print("Warning: openpyxl not available. Excel recording disabled.")

class ExcelHistoryRecorder:
    """Excel历史数据记录器"""
    
    def __init__(self, excel_path: str = None):
        self.excel_path = excel_path
        self.enabled = False
        self._lock = threading.Lock()
        self._write_queue = queue.Queue()
        self._writer_thread = None
        self._stop_writer = threading.Event()
        
        # 缓存已创建的sheet和列映射
        self._sheet_cache = {}
        self._column_cache = {}
        
        if not OPENPYXL_AVAILABLE:
            print("Excel recording not available: openpyxl package missing")
            return
        
        # 如果没有指定路径，先不创建文件，等到enable时再自动生成
        if excel_path and os.path.exists(excel_path):
            self._load_existing_file()
        else:
            self._workbook = None  # 延迟创建
    
    def _load_existing_file(self):
        """加载已存在的Excel文件"""
        try:
            self._workbook = openpyxl.load_workbook(self.excel_path)
            for sheet_name in self._workbook.sheetnames:
                sheet = self._workbook[sheet_name]
                self._sheet_cache[sheet_name] = sheet
                # 读取第一行作为列名
                headers = []
                for cell in sheet[1]:
                    headers.append(cell.value)
                self._column_cache[sheet_name] = headers
            print(f"Loaded existing Excel file: {self.excel_path}")
        except Exception as e:
            print(f"Failed to load existing Excel file {self.excel_path}: {e}")
            self._workbook = None  # 延迟创建
    
    def _create_new_file(self):
        """创建新的Excel文件"""
        try:
            self._workbook = Workbook()
            # 删除默认sheet
            if 'Sheet' in self._workbook.sheetnames:
                self._workbook.remove(self._workbook['Sheet'])
            print(f"Created new Excel workbook")
        except Exception as e:
            print(f"Failed to create new Excel workbook: {e}")
            self._workbook = None
    
    def set_excel_path(self, path: str):
        """设置Excel文件路径"""
        with self._lock:
            if path != self.excel_path:
                self.excel_path = path
                self._sheet_cache.clear()
                self._column_cache.clear()
                if path and os.path.exists(path):
                    self._load_existing_file()
                else:
                    self._workbook = None  # 延迟创建，等到enable时再创建
    
    def enable(self):
        """启用记录"""
        if not OPENPYXL_AVAILABLE:
            print("Cannot enable Excel recording: openpyxl not available")
            return False
        
        # 如果还没有文件路径，自动生成一个
        if not self.excel_path:
            timestamp = time.strftime("%Y%m%d_%H%M%S")
            default_filename = f"SunSpec3_Log_{timestamp}.xlsx"
            try:
                default_dir = os.path.expanduser("~/Documents")
                if not os.path.exists(default_dir):
                    default_dir = os.getcwd()
            except Exception:
                default_dir = os.getcwd()
            self.excel_path = os.path.join(default_dir, default_filename)
        
        # 如果还没有workbook，创建新文件
        if not self._workbook:
            self._create_new_file()
        
        self.enabled = True
        if not self._writer_thread or not self._writer_thread.is_alive():
            self._stop_writer.clear()
            self._writer_thread = threading.Thread(target=self._writer_worker, daemon=True)
            self._writer_thread.start()
        print(f"Excel recording enabled: {self.excel_path}")
        return True
    
    def disable(self):
        """禁用记录"""
        self.enabled = False
        if self._writer_thread and self._writer_thread.is_alive():
            self._stop_writer.set()
            self._writer_thread.join(timeout=2.0)
        print("Excel recording disabled")
    
    def record_model_data(self, model_id: int, parsed_data: Dict[str, Any], timestamp: float = None):
        """记录模型数据"""
        if not self.enabled or not OPENPYXL_AVAILABLE or not self._workbook:
            return
        
        if timestamp is None:
            timestamp = time.time()
        
        # 将数据放入写入队列
        try:
            self._write_queue.put((model_id, parsed_data, timestamp), timeout=1.0)
        except queue.Full:
            print("Excel write queue full, dropping data")
    
    def _ensure_sheet_exists(self, model_id: int, fields: Dict[str, Any]):
        """确保模型对应的sheet存在，并设置列头"""
        sheet_name = f"Model_{model_id}"
        
        if sheet_name not in self._sheet_cache:
            # 创建新sheet
            sheet = self._workbook.create_sheet(title=sheet_name)
            self._sheet_cache[sheet_name] = sheet
            
            # 创建列头：时间戳 + 所有字段的UI显示名（按UI顺序）
            headers = ["Timestamp"]
            # 这里的fields已经是按顺序的了
            field_labels = [data.get('label', name) for name, data in fields.items()]
            headers.extend(field_labels)
            
            # 写入列头
            for col, header in enumerate(headers, 1):
                sheet.cell(row=1, column=col, value=header)
            
            # 更新列缓存为UI Label，后面按label匹配
            self._column_cache[sheet_name] = headers
            
            print(f"Created sheet {sheet_name} with {len(headers)-1} fields")
    
    def _extract_raw_values(self, field_data: Dict[str, Any]) -> Any:
        """从字段数据中提取原始值（不进行符号解析）"""
        if not isinstance(field_data, dict):
            return field_data
        
        # 优先使用raw字段，如果没有则使用value
        if 'raw' in field_data:
            raw_val = field_data['raw']
        elif 'raw_value' in field_data:
            raw_val = field_data['raw_value']
        else:
            raw_val = field_data.get('value')
        
        # 处理数组类型
        if isinstance(raw_val, list):
            # 将数组转换为逗号分隔的字符串
            try:
                return ','.join(str(int(x)) for x in raw_val)
            except (ValueError, TypeError):
                return ','.join(str(x) for x in raw_val)
        
        # 处理字符串类型
        if isinstance(raw_val, str):
            return raw_val
        
        # 处理数值类型
        try:
            return int(raw_val)
        except (ValueError, TypeError):
            return raw_val
    
    def _writer_worker(self):
        """写入线程工作函数"""
        while not self._stop_writer.is_set():
            try:
                # 批量处理写入队列
                batch = []
                while not self._write_queue.empty() and len(batch) < 10:
                    try:
                        item = self._write_queue.get_nowait()
                        batch.append(item)
                    except queue.Empty:
                        break
                
                if not batch:
                    self._stop_writer.wait(0.1)
                    continue
                
                # 处理批量数据
                with self._lock:
                    for model_id, parsed_data, timestamp in batch:
                        self._write_row_to_excel(model_id, parsed_data, timestamp)
                    
                    # 定期保存文件
                    if self.excel_path:
                        try:
                            self._workbook.save(self.excel_path)
                        except Exception as e:
                            print(f"Failed to save Excel file: {e}")
                
            except Exception as e:
                print(f"Excel writer error: {e}")
    
    def _write_row_to_excel(self, model_id: int, parsed_data: Dict[str, Any], timestamp: float):
        """写入一行数据到Excel"""
        try:
            # 确保sheet存在
            self._ensure_sheet_exists(model_id, parsed_data)
            
            sheet_name = f"Model_{model_id}"
            sheet = self._sheet_cache[sheet_name]
            headers = self._column_cache[sheet_name]
            
            # 创建数据行
            row_data = []
            
            # 时间戳
            dt = datetime.fromtimestamp(timestamp)
            row_data.append(dt.strftime("%H:%M:%S"))  # 只保留时分秒
            
            # 按UI Label顺序填充数据
            for label in headers[1:]:  # 跳过时间戳列
                # 在parsed_data中查找该label对应的字段
                found = False
                for field_name, field_data in parsed_data.items():
                    if field_data.get('label') == label:
                        raw_value = self._extract_raw_values(field_data)
                        row_data.append(raw_value)
                        found = True
                        break
                if not found:
                    row_data.append("")  # 未找到对应字段时留空
            
            # 找到下一空行
            next_row = sheet.max_row + 1
            if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
                next_row = 1  # 空sheet
            
            # 写入数据
            for col, value in enumerate(row_data, 1):
                sheet.cell(row=next_row, column=col, value=value)
            
        except Exception as e:
            print(f"Failed to write row for model {model_id}: {e}")
    
    def save_now(self):
        """立即保存文件"""
        if not self._workbook or not self.excel_path:
            return
        
        with self._lock:
            try:
                self._workbook.save(self.excel_path)
                print(f"Excel file saved: {self.excel_path}")
            except Exception as e:
                print(f"Failed to save Excel file: {e}")
    
    def close(self):
        """关闭记录器"""
        self.disable()
        
        if self._workbook and self.excel_path:
            with self._lock:
                try:
                    self._workbook.save(self.excel_path)
                    print(f"Excel file saved on close: {self.excel_path}")
                except Exception as e:
                    print(f"Failed to save Excel file on close: {e}")
        
        if self._workbook:
            self._workbook.close()
            self._workbook = None